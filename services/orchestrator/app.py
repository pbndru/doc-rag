import os
import json
import httpx
import numpy as np
from fastapi import FastAPI, UploadFile, HTTPException, Query
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from sentence_transformers import SentenceTransformer, util
from langchain_openai import ChatOpenAI
from langchain.agents import AgentExecutor, create_openai_functions_agent
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.tools import tool
from langchain.memory import ConversationBufferMemory

app = FastAPI()

# Load sentence transformer model (accurate version)
model = SentenceTransformer('all-mpnet-base-v2')

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://ollama:11434")
LLM_MODEL = os.getenv("LLM_MODEL", "llama3")

# Initialize LLM (Ollama via OpenAI-compatible API)
llm = ChatOpenAI(
    base_url=f"{OLLAMA_URL}/v1",
    api_key="ollama", # Placeholder for Ollama
    model=LLM_MODEL,
    temperature=0
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DOC_PROCESSOR_URL = os.getenv("DOC_PROCESSOR_URL", "http://document-processor:8001")
VECTOR_DB_URL = os.getenv("VECTOR_DB_URL", "http://vector-db:8080")

processed_files = set()

async def check_if_file_exists_in_db(filename: str, client: httpx.AsyncClient):
    """Check if a file with the given title already exists in Weaviate."""
    graphql_query = {
        "query": f'{{ Get {{ Document(where: {{ path: ["title"], operator: Equal, valueText: "{filename}" }}) {{ title }} }} }}'
    }
    try:
        resp = await client.post(f"{VECTOR_DB_URL}/v1/graphql", json=graphql_query)
        resp.raise_for_status()
        data = resp.json()
        results = data.get("data", {}).get("Get", {}).get("Document", [])
        return len(results) > 0
    except Exception as e:
        print(f"Error checking DB for {filename}: {e}", flush=True)
        return False

async def ingest_file(filename: str, client: httpx.AsyncClient):
    """Send a file to the document-processor for ingestion."""
    filepath = os.path.join("/app/documents", filename)
    if not os.path.exists(filepath):
        print(f"File {filepath} not found for ingestion.", flush=True)
        return

    try:
        with open(filepath, "rb") as f:
            files = {"file": (filename, f)}
            resp = await client.post(f"{DOC_PROCESSOR_URL}/ingest", files=files, timeout=60.0)
            resp.raise_for_status()
            print(f"Successfully ingested {filename}", flush=True)
            processed_files.add(filename)
    except Exception as e:
        print(f"Error ingesting {filename}: {e}", flush=True)

import asyncio
import sys

# -------------------------------------------------------------------
# Ensure Weaviate schema exists and start watcher
# -------------------------------------------------------------------
@app.on_event("startup")
async def startup_event():
    print("Orchestrator starting up...", flush=True)
    async with httpx.AsyncClient() as client:
        # Define class Document if not present
        class_obj = {
            "class": "Document",
            "vectorizer": "none",
            "properties": [
                {"name": "title", "dataType": ["text"]},
                {"name": "content", "dataType": ["text"]},
                {"name": "chunkIndex", "dataType": ["int"]}
            ]
        }
        try:
            resp = await client.post(f"{VECTOR_DB_URL}/v1/schema", json=class_obj)
            print(f"Schema creation response: {resp.status_code}", flush=True)
        except Exception as e:
            print(f"Schema creation error: {e}", flush=True)

    # Start watcher
    asyncio.create_task(watcher())
    print("Watcher task created", flush=True)

async def watcher():
    print("Watcher started", flush=True)
    async with httpx.AsyncClient(timeout=60.0) as client:
        # Initial population of processed_files from DB
        docs_path = "/app/documents"
        if os.path.isdir(docs_path):
            for filename in os.listdir(docs_path):
                if filename.lower().endswith(('.pdf', '.docx', '.txt')):
                    if await check_if_file_exists_in_db(filename, client):
                        print(f"File {filename} already in DB, skipping.", flush=True)
                        processed_files.add(filename)

    while True:
        try:
            async with httpx.AsyncClient(timeout=60.0) as client:
                docs_path = "/app/documents" # Explicit path
                if os.path.isdir(docs_path):
                    files = os.listdir(docs_path)
                    for filename in files:
                        if filename.startswith('.'): continue
                        
                        if filename in processed_files:
                            continue
                            
                        if await check_if_file_exists_in_db(filename, client):
                            print(f"File {filename} already in DB, skipping.", flush=True)
                            processed_files.add(filename)
                            continue
                            
                        print(f"Found new file: {filename}. Ingesting...", flush=True)
                        await ingest_file(filename, client)
                else:
                    print(f"Documents directory not found at {docs_path}", flush=True)
        except Exception as e:
            print(f"Watcher error: {e}", flush=True)
        
        await asyncio.sleep(10)

# -------------------------------------------------------------------
# Tools for the Agent
# -------------------------------------------------------------------

@tool
async def search_documents(query: str):
    """Search for relevant document snippets based on a natural language query. 
    Use this for initial discovery and finding specific details.
    """
    async with httpx.AsyncClient() as client:
        # Compute query vector
        embedding = model.encode([query])[0].tolist()
        
        graphql_query = {
            "query": """
            {
              Get {
                Document (
                  hybrid: {
                    query: %s,
                    vector: %s,
                    alpha: 0.5
                  }
                  limit: 5
                ) {
                  title
                  content
                  _additional { score }
                }
              }
            }
            """ % (json.dumps(query), json.dumps(embedding))
        }
        
        resp = await client.post(f"{VECTOR_DB_URL}/v1/graphql", json=graphql_query)
        resp.raise_for_status()
        data = resp.json()
        hits = data.get("data", {}).get("Get", {}).get("Document", [])
        
        if not hits:
            return "No relevant information found."
        
        results = []
        for hit in hits:
            results.append(f"Source: {hit['title']}\nContent: {hit['content']}")
        
        return "\n---\n".join(results)

@tool
async def read_document_content(filename: str):
    """Retrieve the full text content of a specific document by its filename.
    Use this when a search result mentions another document or when you need more context from a known file.
    """
    filepath = os.path.join("documents", filename)
    if not os.path.exists(filepath):
        return f"Error: Document '{filename}' not found."

    if filename.lower().endswith('.txt'):
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif filename.lower().endswith('.pdf'):
        try:
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                return "".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as e:
            return f"Error reading PDF {filename}: {e}"
    elif filename.lower().endswith('.docx'):
        try:
            import docx
            doc = docx.Document(filepath)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            return f"Error reading DOCX {filename}: {e}"
    else:
        return f"Unsupported file type for {filename}"

tools = [search_documents, read_document_content]

# -------------------------------------------------------------------
# Agent Configuration
# -------------------------------------------------------------------

from langchain.agents import AgentExecutor, create_react_agent
from langchain_core.prompts import PromptTemplate

# -------------------------------------------------------------------
# Agent Configuration
# -------------------------------------------------------------------

def get_available_docs():
    docs_path = "/app/documents"
    if os.path.isdir(docs_path):
        return ", ".join([f for f in os.listdir(docs_path) if f.lower().endswith(('.pdf', '.docx', '.txt'))])
    return "No documents available."

template = """Answer the following questions as best you can. You have access to the following tools:

{tools}

Use the following format:

Question: the input question you must answer
Thought: you should always think about what to do
Action: the action to take, should be one of [{tool_names}]
Action Input: the input to the action
Observation: the result of the action
... (this Thought/Action/Action Input/Observation can repeat N times)
Thought: I now know the final answer
Final Answer: the final answer to the original input question

STRICT RULES:
1. ALWAYS use the search_documents tool first.
2. NEVER mention tool names or tool syntax in your Final Answer.
3. Mention the exact filename (e.g. club-lloyds-benefits.pdf) in your Final Answer.

Available Documents: {available_docs}

Question: {input}
Thought: {agent_scratchpad}"""

prompt = PromptTemplate.from_template(template).partial(available_docs=get_available_docs())

agent = create_react_agent(llm, tools, prompt)
agent_executor = AgentExecutor(agent=agent, tools=tools, verbose=True, return_intermediate_steps=True, handle_parsing_errors=True)
# -------------------------------------------------------------------
# Query – search the vector store (Weaviate)
# -------------------------------------------------------------------
@app.post("/query")
async def query_documents(query: dict):
    user_query = query.get("query")
    if not user_query:
        raise HTTPException(status_code=400, detail="Query string missing")
    
    try:
        # Run the agent
        result = await agent_executor.ainvoke({
            "input": user_query,
            "chat_history": []
        })
        
        answer = result.get("output", "I'm sorry, I couldn't process that request.")
        intermediate_steps = result.get("intermediate_steps", [])
        
        # Clean answer: strip any hallucinated tool calls in backticks
        import re
        answer = re.sub(r'`(search_documents|read_document_content)\(.*?\)`', '', answer).strip()
        
        citations = []
        seen_filenames = set()
        filename_to_snippet = {} # Track the best snippet found for each file

        # 1. Extract from tool observations
        for action, observation in intermediate_steps:
            if action.tool == "search_documents":
                parts = observation.split("\n---\n")
                for part in parts:
                    if "Source: " in part:
                        lines = [l for l in part.split("\n") if l.strip()]
                        filename = lines[0].replace("Source: ", "").strip()
                        content = "\n".join(lines[1:]).replace("Content: ", "").strip()
                        
                        # Store the first (usually most relevant) snippet for this file
                        if filename not in filename_to_snippet:
                            filename_to_snippet[filename] = content[:300]
            
            elif action.tool == "read_document_content":
                filename = action.tool_input
                if isinstance(filename, dict):
                    filename = filename.get("filename", str(filename))
                
                if filename not in filename_to_snippet:
                    filename_to_snippet[filename] = observation[:300]

        # 2. Match citations referenced in the final answer
        docs_dir = "/app/documents"
        if os.path.isdir(docs_dir):
            for filename in os.listdir(docs_dir):
                # If filename is mentioned in answer, we want a citation button
                if filename in answer:
                    snippet = filename_to_snippet.get(filename, "Referenced in document.")
                    citations.append({
                        "filename": filename,
                        "snippet": snippet
                    })
                    seen_filenames.add(filename)

        # 3. Add any files the agent searched but didn't explicitly name in the answer (helpful context)
        for filename, snippet in filename_to_snippet.items():
            if filename not in seen_filenames:
                citations.append({
                    "filename": filename,
                    "snippet": snippet
                })
                seen_filenames.add(filename)

        return {"answer": answer, "citations": citations}
        
    except Exception as e:
        print(f"Agent execution error: {e}", flush=True)
        raise HTTPException(status_code=500, detail=str(e))

# -------------------------------------------------------------------
# Serve static documents and document content with highlighting
# -------------------------------------------------------------------
app.mount("/documents", StaticFiles(directory="documents"), name="documents")

@app.get("/document/{filename}")
async def get_document(filename: str, highlight: str = Query(None)):
    """Return document content with optional highlight snippet."""
    filepath = os.path.join("documents", filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")

    if filename.lower().endswith('.txt'):
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return {"filename": filename, "content": content, "highlight": highlight, "type": "txt"}
    elif filename.lower().endswith('.pdf'):
        try:
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                text = "".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as e:
            print(f"Error retrieving document {filename}: {e}")
            raise HTTPException(status_code=500, detail=f"PDF parsing error: {e}")
        return {"filename": filename, "content": text, "highlight": highlight, "type": "pdf"}
    else:
        return {"filename": filename, "highlight": highlight, "type": "unknown"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
